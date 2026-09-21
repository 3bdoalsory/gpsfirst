from flask import Flask, render_template, request, redirect, session, flash, jsonify, send_file
from werkzeug.security import check_password_hash, generate_password_hash
from pathlib import Path
from datetime import date, timedelta, datetime
from functools import wraps
import sqlite3, os, json, math, time, io
from database import init as init_database

app = Flask(__name__)
app.secret_key = os.getenv("GPS_SECRET_KEY", "dev-change-me")
GOOGLE_MAPS_API_KEY = os.getenv("GOOGLE_MAPS_API_KEY", "").strip()
MAP_PROVIDER = os.getenv("MAP_PROVIDER", "leaflet").strip().lower()
DB = Path(os.getenv("GPS_DB_PATH", str(Path(__file__).with_name("gpsplatform.db"))))
init_database()
_LAST_GPS_CLEANUP = 0.0

@app.after_request
def cache_static_assets(response):
    if request.path.startswith('/static/'):
        response.headers['Cache-Control']='public, max-age=604800, immutable'
    return response

def db():
    c = sqlite3.connect(DB)
    c.row_factory = sqlite3.Row
    c.execute("PRAGMA foreign_keys=ON")
    return c



def setting_float(c, key, default):
    try:
        r=c.execute("SELECT value FROM system_settings WHERE key=?",(key,)).fetchone()
        return float(r["value"]) if r else float(default)
    except Exception:
        return float(default)

def stop_thresholds(c):
    return (setting_float(c,"stop_min_seconds",20),
            setting_float(c,"stop_max_speed_kmh",1),
            setting_float(c,"stop_max_drift_m",25))

def point_in_polygon(lat,lon,poly):
    inside=False; j=len(poly)-1
    for i in range(len(poly)):
        yi,xi=poly[i]; yj,xj=poly[j]
        hit=((xi>lon)!=(xj>lon)) and (lat < (yj-yi)*(lon-xi)/((xj-xi) or 1e-12)+yi)
        if hit: inside=not inside
        j=i
    return inside

def parse_device_metrics(raw):
    """SinoTrack V8: third field from end = GSM 0..31; last field = battery %."""
    try:
        parts=str(raw or '').strip().rstrip('#').split(',')
        if len(parts) < 4: return None, None
        gsm=int(float(parts[-3])); battery=int(float(parts[-1]))
        return max(0,min(31,gsm)), max(0,min(100,battery))
    except Exception:
        return None, None

def gsm_status(value):
    try:
        v=int(value)
        return 'strong' if v>15 else ('medium' if v>=7 else 'weak')
    except Exception:
        return 'weak'

def h02_command(device_id, cut=True):
    return f"*HQ,{device_id},S20,{datetime.utcnow().strftime('%H%M%S')},1,{1 if cut else 0}#"

def queue_device_command(c, d, request_id, command_type, cut):
    cmd=h02_command(d['device_id'],cut)
    c.execute("INSERT INTO device_commands(device_pk,request_id,command_type,command_text,status) VALUES(?,?,?,?,?)",
              (d['id'],request_id,command_type,cmd,'queued'))
    return cmd

def signal_status(last_update, speed=0, timeout_seconds=120):
    if not last_update: return "offline"
    try:
        dt=datetime.fromisoformat(str(last_update).replace("Z","+00:00").replace("+00:00",""))
        if (datetime.utcnow()-dt).total_seconds() > timeout_seconds: return "offline"
    except Exception:
        return "offline"
    return "moving" if float(speed or 0) > 1 else "stopped"

def sync_subscription_notifications(c, user_id):
    rows=c.execute("SELECT id,platform_id,name,subscription_end FROM devices WHERE user_id=? AND service_status!='final' AND subscription_end IS NOT NULL",(user_id,)).fetchall()
    today=date.today()
    for d in rows:
        try:
            end=date.fromisoformat(d["subscription_end"]); days=(end-today).days
        except Exception: continue
        if 0 <= days <= 30:
            kind=f'subscription:{d["id"]}:{end.isoformat()}'
            if not c.execute("SELECT 1 FROM notifications WHERE user_id=? AND kind=?",(user_id,kind)).fetchone():
                c.execute("INSERT INTO notifications(user_id,device_pk,kind,title,message) VALUES(?,?,?,?,?)",(user_id,d["id"],kind,"اقتراب انتهاء الاشتراك",f'اشتراك المركبة {d["platform_id"]} · {d["name"]} ينتهي خلال {days} يوم'))

def process_geofences(c,d,lat,lon):
    if not d["user_id"]: return
    rows=c.execute("""SELECT g.* FROM geofences g JOIN geofence_devices gd ON gd.geofence_id=g.id
                      WHERE gd.device_pk=? AND g.user_id=? AND g.is_active=1""",(d["id"],d["user_id"])).fetchall()
    for f in rows:
        try: poly=json.loads(f["polygon_json"])
        except Exception: poly=[]
        if len(poly)<3: continue
        inside=point_in_polygon(lat,lon,poly)
        kind=f'geofence:{f["id"]}:{"in" if inside else "out"}'
        last=c.execute("""SELECT kind FROM notifications WHERE user_id=? AND device_pk=? AND kind LIKE ?
                          ORDER BY id DESC LIMIT 1""",(d["user_id"],d["id"],f'geofence:{f["id"]}:%')).fetchone()
        if last and last["kind"]==kind: continue
        ok=(inside and f["alert_type"] in ("enter","both")) or ((not inside) and f["alert_type"] in ("exit","both"))
        if ok:
            c.execute("""INSERT INTO notifications(user_id,device_pk,kind,title,message) VALUES(?,?,?,?,?)""",
                      (d["user_id"],d["id"],kind,"تنبيه منطقة",f'المركبة {d["platform_id"]} {"دخلت" if inside else "خرجت من"} منطقة {f["name"]}'))

def refresh_pending_immobilize(c,d):
    r=c.execute("SELECT * FROM immobilize_requests WHERE device_pk=? AND status='pending_stop' ORDER BY id DESC LIMIT 1",(d["id"],)).fetchone()
    if not r: return
    if confirmed_stopped(c,d["device_id"]):
        c.execute("UPDATE immobilize_requests SET status='queued',ready_at=CURRENT_TIMESTAMP WHERE id=?",(r["id"],))
        queue_device_command(c,d,r["id"],"immobilize",True)
        c.execute("UPDATE service_audit SET result='queued' WHERE request_id=?",(r["id"],))

def cleanup_old_gps(c):
    global _LAST_GPS_CLEANUP
    now=time.time()
    if now-_LAST_GPS_CLEANUP < 3600: return
    days=max(1,int(setting_float(c,"history_retention_days",90)))
    c.execute(f"DELETE FROM gps_data WHERE created_at < datetime('now','-{days} days')")
    _LAST_GPS_CLEANUP=now

def parse_acc(value):
    if value is None or value=="": return None
    if isinstance(value,bool): return 1 if value else 0
    txt=str(value).strip().lower()
    if txt in ("1","true","on","yes","acc_on"): return 1
    if txt in ("0","false","off","no","acc_off"): return 0
    return None

def parse_h02_raw(raw):
    p=raw.strip().split(",")
    if len(p)<11 or not p[0].startswith("*HQ") or p[4]!="A": raise ValueError("invalid packet")
    def dec(x,d,lon=False):
        n=3 if lon else 2; v=float(x[:n])+float(x[n:])/60
        return -v if d in ("S","W") else v
    return {"device_id":p[1],"latitude":dec(p[5],p[6]),"longitude":dec(p[7],p[8],True),
            "speed":round(float(p[9] or 0)*1.852,2),"heading":float(p[10] or 0),"speed_unit":"kmh","raw_data":raw}

def login_required(admin=False):
    def deco(fn):
        @wraps(fn)
        def wrap(*a, **kw):
            if "user_id" not in session:
                return redirect("/")
            if admin and session.get("role") != "admin":
                return redirect("/dashboard")
            return fn(*a, **kw)
        return wrap
    return deco

def device_state(d):
    if d["service_status"] == "final":
        return "final"
    if d["service_status"] == "temporary":
        return "temporary"
    if d["subscription_end"]:
        try:
            end = date.fromisoformat(d["subscription_end"])
            if end < date.today():
                return "expired"
            if end <= date.today() + timedelta(days=7):
                return "expiring"
        except ValueError:
            pass
    return "active"

def parse_dmy(value):
    value=(value or "").strip()
    if not value: return None
    for fmt in ("%d/%m/%Y","%Y-%m-%d"):
        try: return datetime.strptime(value,fmt).date().isoformat()
        except ValueError: pass
    return None

@app.template_filter("dmy")
def dmy(value):
    if not value: return "-"
    try:
        return datetime.fromisoformat(str(value).replace("Z","")).strftime("%d/%m/%Y")
    except Exception: return str(value)

@app.template_filter("dmytime")
def dmytime(value):
    if not value: return "-"
    try:
        return datetime.fromisoformat(str(value).replace("Z","")).strftime("%d/%m/%Y %H:%M")
    except Exception: return str(value)

def immobilize_allowed_for_current_user(c):
    if session.get("role") == "admin": return True
    u=c.execute("SELECT allow_immobilize FROM users WHERE id=?",(session.get("user_id"),)).fetchone()
    return bool(u and u["allow_immobilize"])

def confirmed_stopped(c, device_id):
    rows=c.execute("SELECT latitude,longitude,speed,acc,created_at FROM gps_data WHERE device_id=? ORDER BY id DESC LIMIT 3",(device_id,)).fetchall()
    if len(rows)<2: return False
    try:
        min_seconds,max_speed,max_drift=stop_thresholds(c)
        newest=datetime.fromisoformat(rows[0]["created_at"]); oldest=datetime.fromisoformat(rows[-1]["created_at"])
        if (newest-oldest).total_seconds()<min_seconds: return False
        if any((r["speed"] or 0)>max_speed for r in rows): return False
        known_acc=[r["acc"] for r in rows if r["acc"] is not None]
        if known_acc and any(int(v)==1 for v in known_acc): return False
        def hav(a,b):
            lat1,lon1,lat2,lon2=map(math.radians,[a["latitude"],a["longitude"],b["latitude"],b["longitude"]])
            h=math.sin((lat2-lat1)/2)**2+math.cos(lat1)*math.cos(lat2)*math.sin((lon2-lon1)/2)**2
            return 6371008.8*2*math.asin(min(1,math.sqrt(h)))
        return all(r["latitude"] is not None and r["longitude"] is not None for r in rows) and hav(rows[0],rows[-1])<=max_drift
    except Exception: return False

def can_access_device(d):
    if not d:
        return False
    return session.get("role") == "admin" or d["user_id"] == session.get("user_id")

@app.route("/", methods=["GET","POST"])
def login():
    if request.method == "POST":
        c=db()
        u=c.execute("SELECT * FROM users WHERE username=?", (request.form["username"].strip(),)).fetchone()
        c.close()
        if not u or not check_password_hash(u["password_hash"], request.form["password"]):
            flash("اسم المستخدم أو كلمة المرور غير صحيحة", "error")
            return render_template("login.html")
        if not u["is_active"]:
            flash("هذا الحساب موقوف من الإدارة", "error")
            return render_template("login.html")
        session.clear()
        session.update(user_id=u["id"], username=u["username"], role=u["role"])
        # "تذكرني" يجعل جلسة تسجيل الدخول مستمرة بدل أن تنتهي مع إغلاق المتصفح.
        session.permanent = bool(request.form.get("remember"))
        if session.permanent:
            app.permanent_session_lifetime = timedelta(days=30)
        return redirect("/admin" if u["role"]=="admin" else "/dashboard")
    if session.get("user_id"):
        return redirect("/admin" if session.get("role")=="admin" else "/dashboard")
    return render_template("login.html")

@app.get("/logout")
def logout():
    session.clear()
    return redirect("/")

@app.get("/dashboard")
@login_required()
def dashboard():
    if session["role"]=="admin":
        return redirect("/admin")
    c=db()
    rows=c.execute("""
      SELECT d.*,
       (SELECT latitude FROM gps_data g WHERE g.device_id=d.device_id ORDER BY g.id DESC LIMIT 1) latitude,
       (SELECT longitude FROM gps_data g WHERE g.device_id=d.device_id ORDER BY g.id DESC LIMIT 1) longitude,
       (SELECT speed FROM gps_data g WHERE g.device_id=d.device_id ORDER BY g.id DESC LIMIT 1) speed,
       (SELECT heading FROM gps_data g WHERE g.device_id=d.device_id ORDER BY g.id DESC LIMIT 1) heading,
       (SELECT created_at FROM gps_data g WHERE g.device_id=d.device_id ORDER BY g.id DESC LIMIT 1) last_update,
       (SELECT gsm_signal FROM gps_data g WHERE g.device_id=d.device_id ORDER BY g.id DESC LIMIT 1) gsm_signal,
       (SELECT battery_percent FROM gps_data g WHERE g.device_id=d.device_id ORDER BY g.id DESC LIMIT 1) battery_percent
      FROM devices d
      WHERE d.user_id=? AND d.service_status!='final'
      ORDER BY d.platform_id
    """,(session["user_id"],)).fetchall()
    sync_subscription_notifications(c, session["user_id"]); c.commit()
    notes=c.execute("SELECT * FROM notifications WHERE user_id=? ORDER BY id DESC",
                    (session["user_id"],)).fetchall()
    unread_count=sum(1 for n in notes if not n["is_read"])
    geo_notes=[n for n in notes if not n["is_read"]]
    audits=c.execute("""SELECT a.*,d.platform_id,d.name device_name
                        FROM service_audit a JOIN devices d ON d.id=a.device_pk
                        WHERE a.user_id=? ORDER BY a.id DESC LIMIT 20""",
                     (session["user_id"],)).fetchall()
    fences=c.execute("""SELECT g.id,g.name,g.polygon_json,g.alert_type,g.is_active,g.color
                        FROM geofences g
                        WHERE g.user_id=? AND g.is_active=1 ORDER BY g.id""",(session["user_id"],)).fetchall()
    c.close()
    devices=[]
    for x in rows:
        z=dict(x); z["state"]=device_state(x); z["subscription_soon"]=False; z["gsm_status"]=gsm_status(x["gsm_signal"])
        if x["subscription_end"]:
            try: z["subscription_soon"]=(date.fromisoformat(x["subscription_end"])-date.today()).days <= 30
            except ValueError: pass
        devices.append(z)
    fence_data=[]
    for f in fences:
        z=dict(f)
        try: z["polygon"]=json.loads(z.get("polygon_json") or "[]")
        except Exception: z["polygon"]=[]
        fence_data.append(z)
    return render_template("dashboard.html", devices=devices, notes=notes, geo_notes=geo_notes, unread_count=unread_count, audits=audits, fences=fence_data)

@app.post("/api/tracker")
def tracker_ingest():
    expected=os.getenv("GPS_INGEST_TOKEN","").strip()
    if expected:
        auth=request.headers.get("Authorization","")
        supplied=auth[7:].strip() if auth.lower().startswith("bearer ") else request.headers.get("X-GPS-Token","").strip()
        if supplied != expected:
            return jsonify(ok=False,error="unauthorized"),401
    payload=request.get_json(silent=True)
    if not isinstance(payload,dict):
        if request.form:
            payload=request.form.to_dict()
        else:
            raw=request.get_data(as_text=True).strip()
            try: payload=parse_h02_raw(raw)
            except Exception: return jsonify(ok=False,error="invalid_payload"),400
    try:
        did=str(payload.get("device_id") or payload.get("id") or payload.get("imei") or "").strip()
        lat=float(payload.get("latitude",payload.get("lat")))
        lon=float(payload.get("longitude",payload.get("lon",payload.get("lng"))))
        speed=float(payload.get("speed",0) or 0)
        unit=str(payload.get("speed_unit",payload.get("unit","kmh"))).strip().lower()
        if unit in ("knot","knots","kt","kts"): speed*=1.852
        heading=float(payload.get("heading",payload.get("course",0)) or 0)
        acc=parse_acc(payload.get("acc"))
        raw_data=payload.get("raw_data",payload.get("raw",json.dumps(payload,ensure_ascii=False)))
        gsm=payload.get("gsm_signal"); battery=payload.get("battery_percent")
        if gsm is None or battery is None:
            rg,rb=parse_device_metrics(raw_data)
            gsm=rg if gsm is None else gsm; battery=rb if battery is None else battery
        gsm=int(gsm) if gsm is not None else None; battery=int(battery) if battery is not None else None
        if not did or not (-90<=lat<=90) or not (-180<=lon<=180): raise ValueError
    except Exception:
        return jsonify(ok=False,error="invalid_payload"),400
    c=db(); d=c.execute("SELECT * FROM devices WHERE device_id=?",(did,)).fetchone()
    if not d:
        c.close(); return jsonify(ok=False,error="unknown_device"),404
    st=device_state(d)
    if st in ("expired","temporary","final") or not d["user_id"]:
        c.close(); return jsonify(ok=False,error="device_unavailable",state=st),409
    c.execute("INSERT INTO gps_data(device_id,latitude,longitude,speed,heading,acc,gsm_signal,battery_percent,raw_data) VALUES(?,?,?,?,?,?,?,?,?)",
              (did,lat,lon,round(speed,2),heading,acc,gsm,battery,str(raw_data)))
    refresh_pending_immobilize(c,d)
    process_geofences(c,d,lat,lon)
    cleanup_old_gps(c)
    c.commit(); c.close()
    return jsonify(ok=True)

@app.get("/api/latest/<int:pid>")
@login_required()
def latest(pid):
    c=db(); d=c.execute("SELECT * FROM devices WHERE platform_id=?",(pid,)).fetchone()
    if not can_access_device(d): c.close(); return jsonify(error="forbidden"),403
    st=device_state(d)
    if st in ("expired","temporary","final"): c.close(); return jsonify(available=False,state=st)
    p=c.execute("""SELECT latitude,longitude,speed,heading,acc,gsm_signal,battery_percent,created_at FROM gps_data
                   WHERE device_id=? ORDER BY id DESC LIMIT 1""",(d["device_id"],)).fetchone()
    c.close()
    payload=dict(p) if p else {}
    return jsonify(available=bool(p),state=st,gsm_status=gsm_status(payload.get("gsm_signal")),**payload)

@app.get("/history/<int:pid>")
@login_required()
def history(pid):
    c=db(); d=c.execute("SELECT * FROM devices WHERE platform_id=?",(pid,)).fetchone(); c.close()
    if not can_access_device(d) or device_state(d) in ("expired","temporary","final"):
        return redirect("/dashboard")
    return render_template("history.html", device=d)

@app.get("/api/history/<int:pid>")
@login_required()
def history_api(pid):
    try:
        s=datetime.fromisoformat(request.args["start"])
        e=datetime.fromisoformat(request.args["end"])
    except:
        return jsonify(error="حدد الفترة"),400
    if e<s or e-s>timedelta(days=90):
        return jsonify(error="الحد الأقصى 90 يومًا"),400
    c=db(); d=c.execute("SELECT * FROM devices WHERE platform_id=?",(pid,)).fetchone()
    if not can_access_device(d) or device_state(d) in ("expired","temporary","final"):
        c.close(); return jsonify(error="forbidden"),403
    rows=c.execute("""SELECT latitude,longitude,speed,heading,acc,created_at FROM gps_data
                      WHERE device_id=? AND created_at BETWEEN ? AND ?
                      ORDER BY created_at""",(d["device_id"],request.args["start"],request.args["end"])).fetchall()
    c.close()
    pts=[dict(r) for r in rows if r["latitude"] is not None and r["longitude"] is not None]
    stops=[]; start=None
    for i,p in enumerate(pts):
        if (p["speed"] or 0)<=1 and start is None: start=i
        if start is not None and ((p["speed"] or 0)>1 or i==len(pts)-1):
            j=i-1 if (p["speed"] or 0)>1 else i
            try:
                t1=datetime.fromisoformat(pts[start]["created_at"]); t2=datetime.fromisoformat(pts[j]["created_at"])
                mins=int((t2-t1).total_seconds()/60)
                if mins>=5:
                    stops.append({"latitude":pts[start]["latitude"],"longitude":pts[start]["longitude"],
                                  "from":pts[start]["created_at"],"to":pts[j]["created_at"],"minutes":mins})
            except: pass
            start=None
    distance_km=0.0
    for a,b in zip(pts,pts[1:]):
        try:
            lat1,lon1,lat2,lon2=map(math.radians,[a["latitude"],a["longitude"],b["latitude"],b["longitude"]])
            h=math.sin((lat2-lat1)/2)**2+math.cos(lat1)*math.cos(lat2)*math.sin((lon2-lon1)/2)**2
            distance_km += 6371.0088 * 2 * math.asin(min(1,math.sqrt(h)))
        except Exception: pass
    max_speed=max([(p["speed"] or 0) for p in pts],default=0)
    return jsonify(points=pts,stops=stops,distance_km=round(distance_km,2),max_speed=round(max_speed,1))

@app.route("/geofences",methods=["GET","POST"])
@login_required()
def geofences():
    if session["role"]=="admin": return redirect("/admin")
    c=db()
    if request.method=="POST":
        try: poly=json.loads(request.form.get("polygon_json","[]"))
        except Exception: poly=[]
        selected=[int(x) for x in request.form.getlist("device_pks") if x.isdigit()]
        allowed={r["id"] for r in c.execute("SELECT id FROM devices WHERE user_id=? AND service_status!='final'",(session["user_id"],)).fetchall()}
        selected=[x for x in selected if x in allowed]
        if len(poly)>=3 and selected:
            cur=c.execute("""INSERT INTO geofences(user_id,device_pk,name,polygon_json,alert_type,sms_enabled,color)\n                             VALUES(?,?,?,?,?,?,?)""",
                          (session["user_id"],selected[0],request.form["name"],json.dumps(poly),request.form["alert_type"],1 if request.form.get("sms_enabled") else 0,request.form.get("color") or "#29c7e8"))
            fid=cur.lastrowid
            c.executemany("INSERT INTO geofence_devices(geofence_id,device_pk) VALUES(?,?)",[(fid,x) for x in selected])
            c.commit(); flash("تم حفظ الزون بنجاح","ok")
        else: flash("ارسم الزون وحدد مركبة واحدة على الأقل","error")
        c.close(); return redirect("/geofences")
    ds=c.execute("SELECT * FROM devices WHERE user_id=? AND service_status!='final' ORDER BY platform_id",(session["user_id"],)).fetchall()
    fs=c.execute("SELECT * FROM geofences WHERE user_id=? ORDER BY id DESC",(session["user_id"],)).fetchall()
    fence_data=[]
    for row in fs:
        z=dict(row)
        try: z["polygon"]=json.loads(z["polygon_json"])
        except Exception: z["polygon"]=[]
        z["device_pks"]=[r["device_pk"] for r in c.execute("SELECT device_pk FROM geofence_devices WHERE geofence_id=?",(z["id"],)).fetchall()]
        fence_data.append(z)
    c.close()
    return render_template("geofences.html",devices=ds,fences=fence_data)

@app.post("/geofences/<int:fid>/edit")
@login_required()
def geofence_edit(fid):
    c=db(); f=c.execute("SELECT * FROM geofences WHERE id=? AND user_id=?",(fid,session["user_id"])).fetchone()
    if not f: c.close(); return redirect("/geofences")
    try: poly=json.loads(request.form.get("polygon_json","[]"))
    except Exception: poly=[]
    selected=[int(x) for x in request.form.getlist("device_pks") if x.isdigit()]
    allowed={r["id"] for r in c.execute("SELECT id FROM devices WHERE user_id=? AND service_status!='final'",(session["user_id"],)).fetchall()}
    selected=[x for x in selected if x in allowed]
    if len(poly)>=3 and selected:
        c.execute("UPDATE geofences SET device_pk=?,name=?,polygon_json=?,alert_type=?,sms_enabled=?,color=?,is_active=? WHERE id=?",
                  (selected[0],request.form["name"],json.dumps(poly),request.form["alert_type"],1 if request.form.get("sms_enabled") else 0,request.form.get("color") or "#29c7e8",1 if request.form.get("is_active") else 0,fid))
        c.execute("DELETE FROM geofence_devices WHERE geofence_id=?",(fid,))
        c.executemany("INSERT INTO geofence_devices(geofence_id,device_pk) VALUES(?,?)",[(fid,x) for x in selected])
        c.commit(); flash("تم تعديل الزون","ok")
    else: flash("يجب إبقاء مركبة واحدة على الأقل ضمن الزون","error")
    c.close(); return redirect("/geofences")

@app.post("/geofences/<int:fid>/delete")
@login_required()
def geofence_delete(fid):
    c=db(); c.execute("DELETE FROM geofence_devices WHERE geofence_id=?",(fid,)); c.execute("DELETE FROM geofences WHERE id=? AND user_id=?",(fid,session["user_id"])); c.commit(); c.close()
    flash("تم حذف الزون","ok"); return redirect("/geofences")

@app.get("/admin")
@login_required(admin=True)
def admin():
    c=db()
    clients=c.execute("""SELECT u.*,COUNT(d.id) device_count FROM users u
                         LEFT JOIN devices d ON d.user_id=u.id
                         WHERE u.role='client' GROUP BY u.id ORDER BY u.id DESC""").fetchall()
    devices=c.execute("""SELECT d.*,u.username FROM devices d LEFT JOIN users u ON u.id=d.user_id
                         ORDER BY d.platform_id""").fetchall()
    audits=c.execute("""SELECT a.*,d.platform_id,d.name device_name FROM service_audit a
                        JOIN devices d ON d.id=a.device_pk ORDER BY a.id DESC LIMIT 100""").fetchall()
    settings={r["key"]:r["value"] for r in c.execute("SELECT key,value FROM system_settings").fetchall()}
    today=date.today(); expired=[]; expiring=[]
    for row in devices:
        if not row["subscription_end"]: continue
        try: days=(date.fromisoformat(row["subscription_end"])-today).days
        except Exception: continue
        item=dict(row); item["days_left"]=days
        if days < 0: expired.append(item)
        elif days <= 30: expiring.append(item)
    expired.sort(key=lambda x:x.get("subscription_end") or "")
    expiring.sort(key=lambda x:x.get("subscription_end") or "")
    c.close()
    return render_template("admin.html",clients=clients,devices=devices,audits=audits,settings=settings,expired=expired,expiring=expiring)

@app.post("/api/notifications/read")
@login_required()
def notifications_read():
    if session.get("role") == "admin": return jsonify(ok=False),403
    c=db(); c.execute("UPDATE notifications SET is_read=1 WHERE user_id=? AND is_read=0",(session["user_id"],)); c.commit(); c.close()
    return jsonify(ok=True)

@app.get("/admin/devices.xlsx")
@login_required(admin=True)
def export_devices_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    c=db(); rows=c.execute("""SELECT d.platform_id,d.device_id,d.name,d.plate,d.vehicle_model,d.vehicle_color,u.username,d.subscription_start,d.subscription_end,d.service_status,d.created_at FROM devices d LEFT JOIN users u ON u.id=d.user_id ORDER BY d.platform_id""").fetchall(); c.close()
    wb=Workbook(); ws=wb.active; ws.title="الأجهزة والاشتراكات"; ws.sheet_view.rightToLeft=True
    headers=["ID الشركة","Device ID الحقيقي","المركبة","اللوحة","الموديل","اللون","العميل","بداية الاشتراك","نهاية الاشتراك","الحالة","تاريخ الإضافة"]
    ws.append(headers)
    for cell in ws[1]: cell.font=Font(bold=True); cell.alignment=Alignment(horizontal="center")
    for r in rows: ws.append([r[k] or "" for k in ["platform_id","device_id","name","plate","vehicle_model","vehicle_color","username","subscription_start","subscription_end","service_status","created_at"]])
    widths=[14,22,22,16,20,14,20,18,18,16,22]
    for i,w in enumerate(widths,1): ws.column_dimensions[chr(64+i)].width=w
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    bio=io.BytesIO(); wb.save(bio); bio.seek(0)
    return send_file(bio,as_attachment=True,download_name=f"gps_devices_{date.today().isoformat()}.xlsx",mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

def _subscription_export(kind):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment
    c=db(); rows=c.execute("""SELECT d.platform_id,d.device_id,d.name,d.plate,d.vehicle_model,d.vehicle_color,u.username,d.subscription_start,d.subscription_end,d.service_status FROM devices d LEFT JOIN users u ON u.id=d.user_id WHERE d.subscription_end IS NOT NULL AND d.subscription_end!='' ORDER BY d.subscription_end""").fetchall(); c.close()
    today=date.today(); selected=[]
    for r in rows:
        try: days=(date.fromisoformat(r["subscription_end"])-today).days
        except Exception: continue
        if (kind=="expired" and days<0) or (kind=="expiring" and 0<=days<=30): selected.append((r,days))
    wb=Workbook(); ws=wb.active; ws.title="منتهية" if kind=="expired" else "قريبة الانتهاء"; ws.sheet_view.rightToLeft=True
    headers=["ID الشركة","Device ID الحقيقي","المركبة","اللوحة","الموديل","اللون","العميل","بداية الاشتراك","نهاية الاشتراك","الحالة / المتبقي"]
    ws.append(headers)
    for cell in ws[1]: cell.font=Font(bold=True); cell.alignment=Alignment(horizontal="center")
    for r,days in selected: ws.append([r["platform_id"],r["device_id"],r["name"],r["plate"] or "",r["vehicle_model"] or "",r["vehicle_color"] or "",r["username"] or "بدون عميل",r["subscription_start"] or "",r["subscription_end"],"منتهي" if days<0 else f"متبقي {days} يوم"])
    for col,w in zip("ABCDEFGHIJ",[14,22,22,16,20,14,20,18,18,18]): ws.column_dimensions[col].width=w
    ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
    out=io.BytesIO(); wb.save(out); out.seek(0)
    return send_file(out,as_attachment=True,download_name=f"subscriptions_{kind}_{date.today().isoformat()}.xlsx",mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.get("/admin/subscriptions/expired.xlsx")
@login_required(admin=True)
def export_expired_subscriptions(): return _subscription_export("expired")

@app.get("/admin/subscriptions/expiring.xlsx")
@login_required(admin=True)
def export_expiring_subscriptions(): return _subscription_export("expiring")

@app.post("/admin/password")
@login_required(admin=True)
def admin_password():
    current=request.form.get("current_password","")
    new=request.form.get("new_password","")
    confirm=request.form.get("confirm_password","")
    c=db(); u=c.execute("SELECT * FROM users WHERE id=? AND role='admin'",(session["user_id"],)).fetchone()
    if not u or not check_password_hash(u["password_hash"],current):
        c.close(); flash("كلمة المرور الحالية غير صحيحة","error"); return redirect("/admin#security")
    if len(new)<4 or new!=confirm:
        c.close(); flash("تأكد من كلمة المرور الجديدة وتأكيدها","error"); return redirect("/admin#security")
    c.execute("UPDATE users SET password_hash=? WHERE id=?",(generate_password_hash(new),u["id"])); c.commit(); c.close()
    flash("تم تغيير كلمة مرور الإدارة","ok"); return redirect("/admin#security")

@app.post("/admin/user")
@login_required(admin=True)
def add_user():
    c=db()
    try:
        c.execute("INSERT INTO users(username,password_hash,role,phone,allow_immobilize) VALUES(?,?,'client',?,?)",
                  (request.form["username"].strip(),generate_password_hash(request.form["password"]),
                   request.form.get("phone","").strip(),1 if request.form.get("allow_immobilize") else 0))
        c.commit();flash("تم إنشاء الحساب","ok")
    except sqlite3.IntegrityError: flash("اسم المستخدم موجود مسبقًا","error")
    c.close();return redirect("/admin#accounts")

@app.post("/admin/user/<int:uid>/edit")
@login_required(admin=True)
def edit_user(uid):
    c=db()
    c.execute("UPDATE users SET username=?,phone=? WHERE id=? AND role='client'",
              (request.form["username"].strip(),request.form.get("phone","").strip(),uid))
    c.execute("UPDATE users SET allow_immobilize=? WHERE id=? AND role=\'client\'",(1 if request.form.get("allow_immobilize") else 0,uid))
    if request.form.get("password","").strip():
        c.execute("UPDATE users SET password_hash=? WHERE id=?",
                  (generate_password_hash(request.form["password"]),uid))
    c.commit();c.close();flash("تم تعديل الحساب","ok");return redirect("/admin#accounts")

@app.post("/admin/user/<int:uid>/toggle")
@login_required(admin=True)
def toggle_user(uid):
    c=db();u=c.execute("SELECT is_active FROM users WHERE id=? AND role='client'",(uid,)).fetchone()
    if u:c.execute("UPDATE users SET is_active=? WHERE id=?",(0 if u["is_active"] else 1,uid));c.commit()
    c.close();return redirect("/admin#accounts")

@app.post("/admin/device")
@login_required(admin=True)
def add_device():
    c=db();pid=c.execute("SELECT COALESCE(MAX(platform_id),10000)+1 n FROM devices").fetchone()["n"]
    try:
        c.execute("""INSERT INTO devices(platform_id,device_id,name,plate,vehicle_model,vehicle_color,user_id,subscription_start,subscription_end)
                     VALUES(?,?,?,?,?,?,?,?,?)""",
                  (pid,request.form["device_id"].strip(),request.form["name"].strip(),request.form.get("plate","").strip(),
                   request.form.get("vehicle_model","").strip(),request.form.get("vehicle_color","").strip(),
                   int(request.form["user_id"]) if request.form.get("user_id") else None,
                   parse_dmy(request.form.get("subscription_start")),parse_dmy(request.form.get("subscription_end"))))
        c.commit();flash(f"تمت إضافة الجهاز برقم الشركة {pid}","ok")
    except sqlite3.IntegrityError:flash("Device ID موجود مسبقًا","error")
    c.close();return redirect("/admin#devices")

@app.post("/admin/device/<int:pid>/edit")
@login_required(admin=True)
def edit_device(pid):
    c=db()
    c.execute("""UPDATE devices SET name=?,plate=?,vehicle_model=?,vehicle_color=?,user_id=?,subscription_start=?,subscription_end=?
                 WHERE platform_id=?""",
              (request.form["name"].strip(),request.form.get("plate","").strip(),
               request.form.get("vehicle_model","").strip(),request.form.get("vehicle_color","").strip(),
               int(request.form["user_id"]) if request.form.get("user_id") else None,
               parse_dmy(request.form.get("subscription_start")),
               parse_dmy(request.form.get("subscription_end")),pid))
    c.commit();c.close();flash("تم تعديل الجهاز","ok");return redirect("/admin#devices")

@app.post("/admin/device/<int:pid>/temporary")
@login_required(admin=True)
def temporary_device(pid):
    c=db(); d=c.execute("SELECT service_status FROM devices WHERE platform_id=?",(pid,)).fetchone()
    if d:
        c.execute("UPDATE devices SET service_status=? WHERE platform_id=?",
                  ("active" if d["service_status"]=="temporary" else "temporary",pid)); c.commit()
    c.close(); return redirect("/admin#devices")

@app.post("/admin/device/<int:pid>/final")
@login_required(admin=True)
def final_device(pid):
    c=db();d=c.execute("SELECT service_status FROM devices WHERE platform_id=?",(pid,)).fetchone()
    if d:
        c.execute("UPDATE devices SET service_status=? WHERE platform_id=?",
                  ("active" if d["service_status"]=="final" else "final",pid));c.commit()
    c.close();return redirect("/admin#devices")

@app.post("/admin/device/<int:pid>/unlink")
@login_required(admin=True)
def unlink_device(pid):
    c=db(); c.execute("UPDATE devices SET user_id=NULL WHERE platform_id=?",(pid,)); c.commit(); c.close()
    flash("تم فك ربط المركبة عن العميل","ok"); return redirect("/admin#devices")

@app.post("/admin/device/<int:pid>/delete")
@login_required(admin=True)
def delete_device(pid):
    c=db(); d=c.execute("SELECT id,device_id FROM devices WHERE platform_id=?",(pid,)).fetchone()
    if not d: c.close(); flash("المركبة غير موجودة","error"); return redirect("/admin#devices")
    c.execute("UPDATE geofences SET device_pk=NULL WHERE device_pk=?",(d["id"],))
    c.execute("DELETE FROM geofence_devices WHERE device_pk=?",(d["id"],))
    c.execute("DELETE FROM notifications WHERE device_pk=?",(d["id"],))
    c.execute("DELETE FROM immobilize_requests WHERE device_pk=?",(d["id"],))
    c.execute("DELETE FROM service_audit WHERE device_pk=?",(d["id"],))
    c.execute("DELETE FROM gps_data WHERE device_id=?",(d["device_id"],))
    c.execute("DELETE FROM devices WHERE id=?",(d["id"],))
    c.commit(); c.close(); flash("تم حذف المركبة وبياناتها من المنصة","ok"); return redirect("/admin#devices")

@app.post("/admin/settings")
@login_required(admin=True)
def admin_settings():
    vals={
      "stop_min_seconds":(request.form.get("stop_min_seconds"),5,300),
      "stop_max_speed_kmh":(request.form.get("stop_max_speed_kmh"),0,20),
      "stop_max_drift_m":(request.form.get("stop_max_drift_m"),1,500),
      "history_retention_days":(request.form.get("history_retention_days"),1,365),
    }
    c=db()
    try:
        for k,(raw,lo,hi) in vals.items():
            v=float(raw); v=max(lo,min(hi,v))
            c.execute("INSERT INTO system_settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",(k,str(v)))
        c.commit(); flash("تم حفظ إعدادات النظام","ok")
    except Exception:
        flash("تحقق من قيم إعدادات النظام","error")
    c.close(); return redirect("/admin#settings")

@app.post("/service/<int:pid>")
@login_required()
def service_request(pid):
    c=db();d=c.execute("SELECT * FROM devices WHERE platform_id=?",(pid,)).fetchone()
    if not can_access_device(d) or device_state(d) in ("expired","temporary","final"):
        c.close();return jsonify(error="forbidden",message="لا يمكن تنفيذ العملية على هذه المركبة"),403
    if not immobilize_allowed_for_current_user(c):
        c.execute("""INSERT INTO service_audit(user_id,username_snapshot,device_pk,action,result) VALUES(?,?,?,?,?)""",
                  (session["user_id"],session["username"],d["id"],"vehicle_stop_request","permission_denied"))
        c.commit();c.close();return jsonify(error="permission_denied",message="هذه العملية ليست ضمن صلاحيات حسابك. يرجى التواصل مع الإدارة."),403
    existing=c.execute("SELECT id,status FROM immobilize_requests WHERE device_pk=? AND status IN ('pending_stop','ready_for_provider') ORDER BY id DESC LIMIT 1",(d["id"],)).fetchone()
    if existing:
        c.close();return jsonify(ok=True,status=existing["status"],message="يوجد طلب إيقاف قائم لهذه المركبة بالفعل")
    status="queued" if confirmed_stopped(c,d["device_id"]) else "pending_stop"
    ready="CURRENT_TIMESTAMP" if status=="queued" else "NULL"
    c.execute(f"""INSERT INTO immobilize_requests(user_id,username_snapshot,device_pk,status,ready_at) VALUES(?,?,?,?,{ready})""",
              (session["user_id"],session["username"],d["id"],status))
    reqid=c.execute("SELECT last_insert_rowid() id").fetchone()["id"]
    if status=="queued":
        queue_device_command(c,d,reqid,"immobilize",True)
    c.execute("""INSERT INTO service_audit(user_id,username_snapshot,device_pk,action,result,request_id) VALUES(?,?,?,?,?,?)""",
              (session["user_id"],session["username"],d["id"],"vehicle_stop_request",status,reqid))
    c.commit();c.close()
    msg="تم تجهيز أمر الإيقاف وسيتم إرساله للمركبة" if status=="queued" else "تم حفظ الطلب وسيبقى بانتظار توقف المركبة الآمن"
    return jsonify(ok=True,status=status,message=msg)


@app.post("/service/<int:pid>/restore")
@login_required()
def service_restore(pid):
    c=db(); d=c.execute("SELECT * FROM devices WHERE platform_id=?",(pid,)).fetchone()
    if not can_access_device(d) or not immobilize_allowed_for_current_user(c):
        c.close(); return jsonify(error="forbidden",message="لا تملك صلاحية هذه العملية"),403
    c.execute("INSERT INTO immobilize_requests(user_id,username_snapshot,device_pk,status,ready_at) VALUES(?,?,?,?,CURRENT_TIMESTAMP)",
              (session["user_id"],session["username"],d["id"],"queued"))
    reqid=c.execute("SELECT last_insert_rowid() id").fetchone()["id"]
    queue_device_command(c,d,reqid,"restore",False)
    c.execute("INSERT INTO service_audit(user_id,username_snapshot,device_pk,action,result,request_id) VALUES(?,?,?,?,?,?)",
              (session["user_id"],session["username"],d["id"],"vehicle_restore_request","queued",reqid))
    c.commit(); c.close()
    return jsonify(ok=True,status="queued",message="تم تجهيز أمر إعادة التشغيل وسيتم إرساله للمركبة")

def expire_stale_device_commands(c, device_pk=None, timeout_seconds=45):
    params=[]
    where="status='sent' AND sent_at IS NOT NULL AND datetime(sent_at) <= datetime('now', ?)"
    params.append(f'-{int(timeout_seconds)} seconds')
    if device_pk is not None:
        where += " AND device_pk=?"
        params.append(device_pk)
    stale=c.execute(f"SELECT * FROM device_commands WHERE {where}",tuple(params)).fetchall()
    for cmd in stale:
        c.execute("UPDATE device_commands SET status='timed_out', response_text=COALESCE(response_text,'device confirmation timeout') WHERE id=?",(cmd['id'],))
        if cmd['request_id']:
            c.execute("UPDATE immobilize_requests SET status='timed_out',completed_at=CURRENT_TIMESTAMP,result='device confirmation timeout' WHERE id=? AND status NOT IN ('confirmed','cancelled')",(cmd['request_id'],))
            req=c.execute("SELECT * FROM immobilize_requests WHERE id=?",(cmd['request_id'],)).fetchone()
            if req:
                action='vehicle_restore_request' if cmd['command_type']=='restore' else 'vehicle_stop_request'
                c.execute("UPDATE service_audit SET result='timed_out' WHERE request_id=?",(req['id'],))
    return len(stale)

def gateway_authorized():
    expected=os.getenv("GPS_GATEWAY_TOKEN",os.getenv("GPS_INGEST_TOKEN","")).strip()
    if not expected: return True
    auth=request.headers.get("Authorization","")
    supplied=auth[7:].strip() if auth.lower().startswith("bearer ") else request.headers.get("X-GPS-Token","").strip()
    return supplied==expected

@app.get("/api/gateway/commands/<device_id>")
def gateway_commands(device_id):
    if not gateway_authorized(): return jsonify(error="unauthorized"),401
    c=db(); d=c.execute("SELECT * FROM devices WHERE device_id=?",(device_id,)).fetchone()
    if d:
        expire_stale_device_commands(c,d["id"]); c.commit()
    if not d: c.close(); return jsonify(commands=[])
    rows=c.execute("SELECT * FROM device_commands WHERE device_pk=? AND status='queued' ORDER BY id LIMIT 5",(d["id"],)).fetchall()
    out=[dict(r) for r in rows]; c.close(); return jsonify(commands=out)

@app.post("/api/gateway/commands/<int:cid>/sent")
def gateway_command_sent(cid):
    if not gateway_authorized(): return jsonify(error="unauthorized"),401
    c=db(); c.execute("UPDATE device_commands SET status='sent',sent_at=CURRENT_TIMESTAMP WHERE id=? AND status='queued'",(cid,)); c.commit(); c.close()
    return jsonify(ok=True)

@app.post("/api/gateway/command-response")
def gateway_command_response():
    if not gateway_authorized(): return jsonify(error="unauthorized"),401
    x=request.get_json(silent=True) or {}; did=str(x.get("device_id") or ""); raw=str(x.get("raw") or ""); cid=x.get("command_id")
    c=db(); d=c.execute("SELECT * FROM devices WHERE device_id=?",(did,)).fetchone()
    if not d: c.close(); return jsonify(error="unknown_device"),404
    cmd=c.execute("SELECT * FROM device_commands WHERE id=? AND device_pk=? AND status='sent'",(cid,d["id"])).fetchone() if cid else None
    if not cmd: cmd=c.execute("SELECT * FROM device_commands WHERE device_pk=? AND status='sent' ORDER BY id DESC LIMIT 1",(d["id"],)).fetchone()
    if cmd:
        failed = "ERROR" in raw.upper() or "FAIL" in raw.upper()
        final_status = "failed" if failed else "confirmed"
        c.execute("UPDATE device_commands SET status=?,confirmed_at=CURRENT_TIMESTAMP,response_text=? WHERE id=?",(final_status,raw,cmd["id"]))
        if cmd["request_id"]:
            c.execute("UPDATE immobilize_requests SET status=?,completed_at=CURRENT_TIMESTAMP,result=? WHERE id=?",(final_status,raw,cmd["request_id"]))
            req=c.execute("SELECT * FROM immobilize_requests WHERE id=?",(cmd["request_id"],)).fetchone()
            if req:
                action="vehicle_restore_request" if cmd["command_type"]=="restore" else "vehicle_stop_request"
                c.execute("UPDATE service_audit SET result=? WHERE request_id=?",
                          ("device_failed" if failed else "confirmed_by_device",req["id"]))
    c.commit(); c.close(); return jsonify(ok=True)

@app.post("/api/immobilize/<int:pid>/cancel")
@login_required()
def immobilize_cancel(pid):
    c=db(); d=c.execute("SELECT * FROM devices WHERE platform_id=?",(pid,)).fetchone()
    if not can_access_device(d): c.close(); return jsonify(error="forbidden"),403
    r=c.execute("SELECT * FROM immobilize_requests WHERE device_pk=? AND status IN ('pending_stop','ready_for_provider') ORDER BY id DESC LIMIT 1",(d["id"],)).fetchone()
    if not r: c.close(); return jsonify(ok=False,message="لا يوجد طلب قابل للإلغاء"),404
    c.execute("UPDATE immobilize_requests SET status='cancelled',completed_at=CURRENT_TIMESTAMP,result='cancelled' WHERE id=?",(r["id"],))
    c.execute("UPDATE service_audit SET result='cancelled' WHERE request_id=?",(r["id"],))
    c.commit(); c.close(); return jsonify(ok=True,status="cancelled",message="تم إلغاء الطلب")

@app.get("/api/immobilize/<int:pid>")
@login_required()
def immobilize_status(pid):
    c=db();d=c.execute("SELECT * FROM devices WHERE platform_id=?",(pid,)).fetchone()
    if not can_access_device(d): c.close(); return jsonify(error="forbidden"),403
    expire_stale_device_commands(c,d["id"]); c.commit()
    r=c.execute("SELECT id,status,requested_at,ready_at,completed_at FROM immobilize_requests WHERE device_pk=? ORDER BY id DESC LIMIT 1",(d["id"],)).fetchone()
    cmd=c.execute("SELECT id,command_type,status,created_at,sent_at,confirmed_at FROM device_commands WHERE device_pk=? ORDER BY id DESC LIMIT 1",(d["id"],)).fetchone()
    out={"request":dict(r) if r else None,"command":dict(cmd) if cmd else None}; c.close()
    return jsonify(**out)

if __name__=="__main__":
    app.run(host="0.0.0.0", port=5000, debug=True)
